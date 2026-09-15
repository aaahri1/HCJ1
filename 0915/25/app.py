import os
import tempfile

from flask import Flask, jsonify, render_template, request
from pypdf import PdfReader
from docx import Document

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 20 * 1024 * 1024  # 20MB

ALLOWED_EXTENSIONS = {".pdf", ".docx"}


def normalize_text(text: str) -> str:
    return text.replace("\r\n", "\n").replace("\r", "\n").replace("\t", " ").replace("\n", "")


def build_stats(text: str, image_count: int) -> dict:
    clean_text = normalize_text(text)
    return {
        "charsWithSpace": len(clean_text),
        "charsWithoutSpace": len(clean_text.replace(" ", "")),
        "wordCount": len(clean_text.split()),
        "spaceCount": clean_text.count(" "),
        "imageCount": image_count,
    }


def analyze_pdf(file_path: str) -> dict:
    reader = PdfReader(file_path)

    text_parts = []
    image_count = 0
    for page in reader.pages:
        text_parts.append(page.extract_text() or "")
        try:
            image_count += len(page.images)
        except Exception:
            pass

    return build_stats("\n".join(text_parts), image_count)


def analyze_docx(file_path: str) -> dict:
    document = Document(file_path)

    text_parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                text_parts.append(cell.text)

    image_count = sum(
        1
        for rel in document.part.rels.values()
        if rel.reltype.endswith("/image") or (rel.target_part is not None and rel.target_part.content_type.startswith("image/"))
    )

    return build_stats("\n".join(text_parts), image_count)


@app.route("/")
def index():
    return render_template("index.html")


@app.route("/analyze", methods=["POST"])
def analyze():
    uploaded_file = request.files.get("file")
    if uploaded_file is None or uploaded_file.filename == "":
        return jsonify({"error": "파일을 선택해주세요."}), 400

    filename = uploaded_file.filename
    ext = os.path.splitext(filename)[1].lower()
    if ext not in ALLOWED_EXTENSIONS:
        return jsonify({"error": "PDF 또는 DOCX 파일만 업로드할 수 있습니다."}), 400

    with tempfile.NamedTemporaryFile(delete=False, suffix=ext) as tmp:
        uploaded_file.save(tmp.name)
        tmp_path = tmp.name

    try:
        if ext == ".pdf":
            stats = analyze_pdf(tmp_path)
        else:
            stats = analyze_docx(tmp_path)
    except Exception as exc:
        return jsonify({"error": f"파일을 분석하는 중 오류가 발생했습니다: {exc}"}), 500
    finally:
        os.remove(tmp_path)

    stats["filename"] = filename
    return jsonify(stats)


if __name__ == "__main__":
    app.run(debug=True)
